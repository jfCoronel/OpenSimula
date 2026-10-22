"""ASHRAE 140-2023 Section 7 (building thermal envelope and fabric load tests, LOAD_TEST folder).

Defines the cases to simulate, how the Standard's outputs are extracted from an OpenSimula
project and where the example results of the reference programs are in
Std140_TF_Results.xlsx.
"""

from pathlib import Path

import numpy as np
import openpyxl

NAME = "LOAD_TEST"
TITLE = "Building thermal envelope and fabric load tests (Section 7)"
FOLDER = Path(__file__).resolve().parent.parent / "LOAD_TEST"
RESULTS_FILE = FOLDER / "Std140_TF_Results.xlsx"
REFERENCE_PROGRAMS = ["BSIMAC", "CSE", "DeST", "EnergyPlus", "ESP-r", "TRNSYS"]

# Same order as the Standard's output form (sheet YourData)
LOAD_CASES = [
    "600", "610", "620", "630", "640", "650", "660", "670", "680", "685", "695",
    "900", "910", "920", "930", "940", "950", "960", "980", "985", "995",
    "195", "200", "210", "215", "220", "230", "240", "250", "270", "280", "290",
    "300", "310", "320", "395", "400", "410", "420", "430", "440", "450", "460",
    "470", "800", "810",
]
FREE_FLOAT_CASES = ["600FF", "900FF", "650FF", "950FF", "680FF", "980FF", "960"]

# Case 960 is the two-zone sunspace case: loads are reported for the conditioned back zone
# and free-float temperatures for the unconditioned sun zone.
LOAD_SPACE = {"960": "back_zone"}
FREE_FLOAT_SPACE = {"960": "sun_zone"}
DEFAULT_SPACE = "spaces_1"

# Output name -> (title, unit, column of the value in the Standard's output form).
# Peaks are followed in the form by three columns with their month, day and hour.
LOAD_OUTPUTS = {
    "annual_heating": ("Annual heating load", "MWh", 3),
    "annual_cooling": ("Annual cooling load", "MWh", 4),
    "peak_heating": ("Annual hourly integrated peak heating load", "kW", 5),
    "peak_cooling": ("Annual hourly integrated peak sensible cooling load", "kW", 9),
}
FREE_FLOAT_OUTPUTS = {
    "average_temperature": ("Annual hourly average zone temperature", "°C", 3),
    "minimum_temperature": ("Annual hourly minimum zone temperature", "°C", 4),
    "maximum_temperature": ("Annual hourly maximum zone temperature", "°C", 8),
}
OUTPUTS = {**LOAD_OUTPUTS, **FREE_FLOAT_OUTPUTS}

# Charts of the validation page: (title, output names drawn side by side)
FIGURES = [
    ("Annual loads", ["annual_heating", "annual_cooling"]),
    ("Peak loads", ["peak_heating", "peak_cooling"]),
    (
        "Free-float zone temperatures",
        ["minimum_temperature", "average_temperature", "maximum_temperature"],
    ),
]


def cases():
    """Cases to simulate, each one once even if it reports both loads and temperatures."""
    return list(dict.fromkeys(LOAD_CASES + FREE_FLOAT_CASES))


def case_file(case):
    return FOLDER / "cases" / f"case_{case}.json"


def _date(dates, i):
    """Month, day and hour (1-24, hour ending) of time step i, as the Standard asks for."""
    d = dates[i]
    return f"{d:%b} {d.day} {d.hour + 1}"


def extract(case, project):
    """Standard's outputs of a simulated case: list of (output, value, date) tuples."""
    dates = project.dates()
    hours = project.parameter("time_step").value / 3600
    rows = []
    if case in LOAD_CASES:
        load = project.component(LOAD_SPACE.get(case, DEFAULT_SPACE)).variable("system_sensible_heat").values
        heating = np.where(load > 0, load, 0)
        cooling = np.where(load < 0, -load, 0)
        rows.append(("annual_heating", heating.sum() * hours / 1e6, ""))
        rows.append(("annual_cooling", cooling.sum() * hours / 1e6, ""))
        i_heat, i_cool = np.argmax(heating), np.argmax(cooling)
        rows.append(("peak_heating", heating[i_heat] / 1000, _date(dates, i_heat)))
        rows.append(("peak_cooling", cooling[i_cool] / 1000, _date(dates, i_cool)))
    if case in FREE_FLOAT_CASES:
        temp = project.component(FREE_FLOAT_SPACE.get(case, DEFAULT_SPACE)).variable("temperature").values
        i_min, i_max = np.argmin(temp), np.argmax(temp)
        rows.append(("average_temperature", temp.mean(), ""))
        rows.append(("minimum_temperature", temp[i_min], _date(dates, i_min)))
        rows.append(("maximum_temperature", temp[i_max], _date(dates, i_max)))
    return rows


def _read_block(ws, header_row, outputs):
    """Values of a table of the output form, starting below its 'CASE#' header row."""
    values = {}
    row = header_row + 1
    while (label := ws.cell(row, 2).value) not in (None, ""):
        case = str(label).strip()
        for output, (_, _, col) in outputs.items():
            value = ws.cell(row, col).value
            if isinstance(value, (int, float)):
                values[(case, output)] = float(value)
        row += 1
    return values


def read_reference():
    """Example results of the reference programs: {program: {(case, output): value}}."""
    wb = openpyxl.load_workbook(RESULTS_FILE, data_only=True)
    reference = {}
    for program in REFERENCE_PROGRAMS:
        ws = wb[program]
        headers = [r for r in range(1, 200) if ws.cell(r, 2).value == "CASE#"]
        # First table: conditioned zone loads; second: free-float temperatures
        values = _read_block(ws, headers[0], LOAD_OUTPUTS)
        values.update(_read_block(ws, headers[1], FREE_FLOAT_OUTPUTS))
        reference[program] = values
    return reference
