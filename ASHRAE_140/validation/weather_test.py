"""ASHRAE 140-2023 weather drivers tests WD100–WD600 (WEATHER_TEST folder, results workbook
Std140_WD_Annual_Results.xlsx).

The six cases use the same building (surfaces in every orientation and tilt) with a different
weather file and 10 min time steps. The outputs are annual averages of the weather variables
and of the total, beam and diffuse solar radiation on the horizontal and on nine surfaces.
Each result is identified by (case, output, point) with point "".
"""

from pathlib import Path

import openpyxl

NAME = "WEATHER_TEST"
TITLE = "Weather drivers tests (WD100–WD600)"
FOLDER = Path(__file__).resolve().parent.parent / "WEATHER_TEST"
RESULTS_FILE = FOLDER / "Std140_WD_Annual_Results.xlsx"
SHEET = "Annual Average Results"
# Column of the sheet -> source of the results ("Actual": values of the weather file itself)
REFERENCE_COLUMNS = {3: "Weather file", 4: "BSIMAC", 5: "CSE", 6: "DeST", 7: "EnergyPlus", 8: "ESP-r",
                     9: "IBPSA", 10: "IDA ICE", 11: "TRNSYS"}
REFERENCE_PROGRAMS = list(REFERENCE_COLUMNS.values())

CASES = {"WD100": 7, "WD200": 50, "WD300": 93, "WD400": 136, "WD500": 179, "WD600": 222}  # block row

# (output, title, unit, File_met variable and factor) for the weather rows (block row + 1...)
_WEATHER = [
    ("t_db", "Dry bulb temperature", "°C", "temperature", 1),
    ("rh", "Relative humidity", "%", "rel_humidity", 1),
    ("w", "Humidity ratio", "kg/kg", "abs_humidity", 0.001),
    ("t_dp", "Dew point temperature", "°C", "dew_point_temp", 1),
    ("t_wb", "Wet bulb temperature", "°C", "wet_bulb_temp", 1),
    ("wind_speed", "Wind speed", "m/s", "wind_speed", 1),
    ("wind_direction", "Wind direction", "° from north", "wind_direction", 1),
    ("pressure", "Station pressure", "mbar", "pressure", 0.01),
    ("total_cloud", "Total cloud cover", "tenths", "total_cloud_cover", 0.1),
    ("opaque_cloud", "Opaque cloud cover", "tenths", "opaque_cloud_cover", 0.1),
    ("t_sky", "Sky temperature", "°C", "sky_temperature", 1),
]
# Surfaces of the output form (in order) -> building surface of the case files
SURFACES = [("horizontal", "Horizontal", "roof"), ("south_90", "South 90", "south_wall"),
            ("east_90", "East 90", "east_wall"), ("north_90", "North 90", "north_wall"),
            ("west_90", "West 90", "west_wall"), ("se_90", "45° east of south 90", "south_east_wall"),
            ("sw_90", "45° west of south 90", "south_west_wall"), ("east_30", "East 30", "east_roof"),
            ("south_30", "South 30", "south_roof"), ("west_30", "West 30", "west_roof")]
RADIATION = [("total", "total"), ("beam", "beam"), ("diffuse", "diffuse")]

OUTPUTS = {name: (f"Annual average {title.lower()}", unit) for name, title, unit, *_ in _WEATHER}
for surface, title, _ in SURFACES:
    for kind, kind_title in RADIATION:
        OUTPUTS[f"{surface}_{kind}"] = (f"Annual average {kind_title} radiation, {title}", "W/m²")


def _cells():
    cells = {}
    names = [name for name, *_ in _WEATHER] + [f"{s}_{k}" for s, _, _ in SURFACES for k, _ in RADIATION]
    for case, block in CASES.items():
        for i, name in enumerate(names, 1):
            cells[(case, name, "")] = (block + i, 0)  # column set by the reference source
    return cells


CELLS = _cells()

GROUPS = [
    {"title": "Weather variables", "kind": "annual", "outputs": [n for n, *_ in _WEATHER]},
    {"title": "Total solar radiation", "kind": "annual", "outputs": [f"{s}_total" for s, _, _ in SURFACES]},
    {"title": "Beam solar radiation", "kind": "annual", "outputs": [f"{s}_beam" for s, _, _ in SURFACES]},
    {"title": "Diffuse solar radiation", "kind": "annual", "outputs": [f"{s}_diffuse" for s, _, _ in SURFACES]},
]


def cases():
    return list(CASES)


def case_file(case):
    return FOLDER / "cases" / f"case_{case}.json"


def extract(case, project):
    """Standard's outputs of a simulated case: list of (output, point, value, date) tuples."""
    met = project.parameter("simulation_file_met").component
    rows = [(name, "", float(met.variable(variable).values.mean() * factor), "")
            for name, _, _, variable, factor in _WEATHER]
    for surface, _, component in SURFACES:
        s = project.component(component)
        beam, diffuse = s.variable("E_dir").values, s.variable("E_dif").values
        rows += [(f"{surface}_total", "", float((beam + diffuse).mean()), ""),
                 (f"{surface}_beam", "", float(beam.mean()), ""),
                 (f"{surface}_diffuse", "", float(diffuse.mean()), "")]
    return rows


def read_reference():
    """Results of the reference programs and of the weather file itself:
    {source: {(case, output, point): value}}."""
    ws = openpyxl.load_workbook(RESULTS_FILE, data_only=True)[SHEET]
    for case, block in CASES.items():
        if str(ws.cell(block, 1).value).strip() != case:
            raise ValueError(f"{RESULTS_FILE.name}[{SHEET}] row {block}: expected {case}")
    reference = {}
    for col, source in REFERENCE_COLUMNS.items():
        reference[source] = {
            key: float(value)
            for key, (row, _) in CELLS.items()
            if isinstance(value := ws.cell(row, col).value, (int, float))
        }
    return reference
