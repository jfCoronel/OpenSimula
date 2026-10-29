"""ASHRAE 140-2023 air-side HVAC equipment analytical verification tests AE101–AE445
(AIR_SIDE_HVAC_TEST folder).

All cases are steady state: outputs are for the last hour of the simulation. The output form
has four blocks (fan coil AE100, single zone AE200, constant volume terminal reheat AE300,
variable air volume AE400), each with summary loads per case and the psychrometric state
(T, w, v, h, m) at every location of the air system. The reference results are in two
workbooks with the same layout: AE100/AE200 from Std140_AE_FCSZ_Results.xlsx and
AE300/AE400 from Std140_AE_CVVV_Results.xlsx.

Each result is identified by (case, output, point): point is "" for the summary outputs and
the location code (oa, ma, cco, sa, z1, ...) for the state outputs. Loads are compared as
magnitudes, since the reference programs report them with different sign conventions.
"""

from pathlib import Path

import openpyxl
import psychrolib as sicro

NAME = "AIR_SIDE_HVAC_TEST"
TITLE = "Air-side HVAC equipment tests (AE101–AE445)"
FOLDER = Path(__file__).resolve().parent.parent / "AIR_SIDE_HVAC_TEST"
RESULTS_FILE = FOLDER / "Std140_AE_FCSZ_Results.xlsx"  # shown on the page; see WORKBOOKS
WORKBOOKS = {"Std140_AE_FCSZ_Results.xlsx": ("100", "200"), "Std140_AE_CVVV_Results.xlsx": ("300", "400")}
REFERENCE_SHEETS = {"QAS": "QAS", "DEEAP": "DEEAP", "DeST": "DeST", "DOE22": "DOE-2.2",
                    "Eplus": "EnergyPlus", "IES-VE": "IES-VE", "LCEM": "LCEM", "TRNSYS": "TRNSYS"}
REFERENCE_PROGRAMS = list(REFERENCE_SHEETS.values())

SERIES = {  # series -> (cases, first summary row, state locations from column S)
    "100": (["AE101", "AE103", "AE104"], 61, ["oa", "ma", "hco", "cco", "sa", "z1", "ra"]),
    "200": (["AE201", "AE203", "AE204", "AE205", "AE206", "AE226", "AE245"], 89,
            ["oa", "ma", "hco", "cco", "sa", "z1", "rfi", "rfo"]),
    "300": (["AE301", "AE303", "AE304", "AE305", "AE306", "AE326", "AE345"], 137,
            ["oa", "ma", "pco", "cco", "sa", "z1s", "z2s", "z1", "z2", "rfi", "rfo"]),
    "400": (["AE401", "AE403", "AE404", "AE405", "AE406", "AE426", "AE445"], 185,
            ["oa", "ma", "pco", "cco", "sa", "z1s", "z2s", "z1", "z2", "rfi", "rfo"]),
}
SINGLE_ZONE_SUMMARY = ["qzh", "qzc", "qzl", "qh", "qcs", "qcl", "qct", "rhcco"]  # columns C-J
MULTI_ZONE_SUMMARY = ["qzh1", "qzc1", "qzl1", "qzh2", "qzc2", "qzl2", "qh", "qcs", "qcl", "qct",
                      "rhcco", "qh1", "qh2"]  # columns C-O
STATE_VARIABLES = ["T", "w", "v", "h", "m"]  # 5 rows per case

OUTPUTS = {
    "qzh": ("Zone sensible heating load", "kWh/h"),
    "qzc": ("Zone sensible cooling load", "kWh/h"),
    "qzl": ("Zone latent load", "kWh/h"),
    "qzh1": ("Zone 1 sensible heating load", "kWh/h"),
    "qzc1": ("Zone 1 sensible cooling load", "kWh/h"),
    "qzl1": ("Zone 1 latent load", "kWh/h"),
    "qzh2": ("Zone 2 sensible heating load", "kWh/h"),
    "qzc2": ("Zone 2 sensible cooling load", "kWh/h"),
    "qzl2": ("Zone 2 latent load", "kWh/h"),
    "qh": ("Heating (preheat) coil load", "kWh/h"),
    "qcs": ("Cooling coil sensible load", "kWh/h"),
    "qcl": ("Cooling coil latent load", "kWh/h"),
    "qct": ("Cooling coil total load", "kWh/h"),
    "rhcco": ("Cooling coil outlet relative humidity", "%"),
    "qh1": ("Zone 1 reheat coil load", "kWh/h"),
    "qh2": ("Zone 2 reheat coil load", "kWh/h"),
    "T": ("Dry bulb temperature", "°C"),
    "w": ("Humidity ratio", "kg/kg"),
    "v": ("Specific volume", "L/kg"),
    "h": ("Enthalpy", "kJ/kg"),
    "m": ("Dry air mass flow rate", "kg/s"),
}
LOADS = {o for o, (_, unit) in OUTPUTS.items() if unit == "kWh/h"}


def _cells():
    cells = {}
    for series, (cases, first_row, locations) in SERIES.items():
        summary = SINGLE_ZONE_SUMMARY if series in ("100", "200") else MULTI_ZONE_SUMMARY
        for i, case in enumerate(cases):
            for col, output in enumerate(summary, 3):
                cells[(case, output, "")] = (first_row + i, col)
            for k, variable in enumerate(STATE_VARIABLES):
                for col, location in enumerate(locations, 19):
                    cells[(case, variable, location)] = (first_row + 5 * i + k, col)
    return cells


CELLS = _cells()


def _groups():
    groups = []
    for series, (cases, _, _) in SERIES.items():
        summary = SINGLE_ZONE_SUMMARY if series in ("100", "200") else MULTI_ZONE_SUMMARY
        loads = [o for o in summary if o != "rhcco"]
        groups.append({"title": f"AE{series} series: loads", "kind": "annual", "cases": cases,
                       "outputs": loads})
        groups.append({"title": f"AE{series} series: coil outlet humidity and air states", "kind": "annual",
                       "cases": cases, "outputs": ["rhcco"] + STATE_VARIABLES, "figure": False})
    return groups


GROUPS = _groups()


def cases():
    return [c for cases, _, _ in SERIES.values() for c in cases]


def case_file(case):
    return FOLDER / "cases" / f"case_{case}.json"


def extract(case, project):
    """Standard's outputs of a simulated case: list of (output, point, value, date) tuples."""
    sicro.SetUnitSystem(sicro.SI)
    series = next(s for s, (cases, _, _) in SERIES.items() if case in cases)
    system = project.component("system")
    p = system.props["ATM_PRESSURE"]

    def last(component, variable):
        return float(project.component(component).variable(variable).values[-1])

    def s(variable):
        return float(system.variable(variable).values[-1])

    def state(T, w, m):  # w in kg/kg; m None where not applicable
        return {"T": T, "w": w, "v": sicro.GetMoistAirVolume(T, w, p) * 1000,
                "h": sicro.GetMoistAirEnthalpy(T, w) / 1000, "m": m}

    m_su = s("m_air_flow")
    m_oa = s("outdoor_air_fraction") * m_su
    m_ra = s("m_return_air_flow")  # supply - minimum outdoor air (relieved with the economizer)
    w_oa = sicro.GetHumRatioFromTWetBulb(s("T_OA"), s("T_OAwb"), p)
    w_ma = sicro.GetHumRatioFromTWetBulb(s("T_MA"), s("T_MAwb"), p)
    coil = state(s("T_CA"), s("w_CA") / 1000, m_su)
    supply_w = s("w_SA") / 1000
    states = {"oa": state(s("T_OA"), w_oa, m_oa), "ma": state(s("T_MA"), w_ma, m_su),
              "sa": state(s("T_SA"), supply_w, m_su)}
    q_s, q_l = s("Q_sensible") / 1000, s("Q_latent") / 1000
    cooling = q_s < 0 or q_l > 0
    rows = [("qh", "", max(q_s, 0), ""),
            ("qcs", "", max(-q_s, 0), ""),
            ("qcl", "", q_l if cooling else 0, ""),
            ("qct", "", max(-q_s, 0) + (q_l if cooling else 0), ""),
            ("rhcco", "", sicro.GetRelHumFromHumRatio(coil["T"], coil["w"], p) * 100, "")]

    def zone(space):  # sensible (+ heating) and latent zone loads, kW
        sen = last(space, "system_sensible_heat") + last(space, "u_system_sensible_heat")
        lat = last(space, "system_latent_heat") + last(space, "u_system_latent_heat")
        return sen / 1000, abs(lat) / 1000

    if series in ("100", "200"):
        sen, lat = zone("spaces_1")
        rows += [("qzh", "", max(sen, 0), ""), ("qzc", "", max(-sen, 0), ""), ("qzl", "", lat, "")]
        zone_w = last("spaces_1", "abs_humidity") / 1000
        zone_state = state(last("spaces_1", "temperature"), zone_w, None)
        states.update({"hco": coil, "cco": coil, "z1": zone_state})
        if series == "100":
            states["ra"] = state(zone_state["T"], zone_w, m_ra)
        else:
            states["rfi"] = state(zone_state["T"], zone_w, m_ra)
            states["rfo"] = state(s("T_RA"), zone_w, m_ra)
    else:
        for n, space in [("1", "spaces_1"), ("2", "spaces_2")]:
            sen, lat = zone(space)
            rows += [(f"qzh{n}", "", max(sen, 0), ""), (f"qzc{n}", "", max(-sen, 0), ""),
                     (f"qzl{n}", "", lat, "")]
        rows += [("qh1", "", s("Q_reheat_0") / 1000, ""), ("qh2", "", s("Q_reheat_1") / 1000, "")]
        return_w = s("w_RA") / 1000
        states.update({
            "pco": coil, "cco": coil,
            "z1s": state(s("T_SA_0"), supply_w, s("m_air_flow_0")),
            "z2s": state(s("T_SA_1"), supply_w, s("m_air_flow_1")),
            "z1": state(last("spaces_1", "temperature"), last("spaces_1", "abs_humidity") / 1000, None),
            "z2": state(last("spaces_2", "temperature"), last("spaces_2", "abs_humidity") / 1000, None),
            "rfi": state(s("T_ZA"), return_w, m_ra),
            "rfo": state(s("T_RA"), return_w, m_ra),
        })
    for location in SERIES[series][2]:
        for variable, value in states[location].items():
            if value is not None:
                rows.append((variable, location, float(value), ""))
    return [(o, pt, float(v), d) for o, pt, v, d in rows]


def read_reference():
    """Results of the reference programs: {program: {(case, output, point): value}}.

    Loads are taken as magnitudes (the programs use different sign conventions)."""
    reference = {program: {} for program in REFERENCE_PROGRAMS}
    for workbook, series_list in WORKBOOKS.items():
        wb = openpyxl.load_workbook(FOLDER / workbook, data_only=True)
        series_cases = {c for s in series_list for c in SERIES[s][0]}
        for sheet, program in REFERENCE_SHEETS.items():
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            for s in series_list:  # check the layout: case labels of the summary rows
                cases, first_row, _ = SERIES[s]
                for i, case in enumerate(cases):
                    label = str(ws.cell(first_row + i, 2).value).strip()
                    if label != case:
                        raise ValueError(f"{workbook}[{sheet}] row {first_row + i}: {label}, expected {case}")
            for key, (row, col) in CELLS.items():
                if key[0] in series_cases and isinstance(value := ws.cell(row, col).value, (int, float)):
                    reference[program][key] = abs(float(value)) if key[1] in LOADS else float(value)
    return reference
