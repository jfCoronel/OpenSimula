"""ASHRAE 140-2023 HVAC equipment performance tests, unitary space cooling, steady-state cases
CE100–CE200 (COOLING_TEST folder, results workbook Std140_CE_a_Results.xlsx).

All outputs are for February: totals of consumption and loads, and means, maxima and
minima of COP, indoor dry bulb and zone humidity ratio. Each result is identified by
(case, output, point) with point "".
"""

from pathlib import Path

import numpy as np
import openpyxl

from dx_system import dx_hourly

NAME = "COOLING_TEST_A"
TITLE = "Space cooling equipment, steady-state tests (CE100–CE200)"
FOLDER = Path(__file__).resolve().parent.parent / "COOLING_TEST"
RESULTS_FILE = FOLDER / "Std140_CE_a_Results.xlsx"
# Sheet of the results workbook -> program or analytical solution
REFERENCE_SHEETS = {
    "Analytical-TUD": "Analytical TUD", "Analytical-HTAL1": "Analytical HTAL1",
    "Analytical-HTAL2": "Analytical HTAL2", "CA-SIS": "CA-SIS", "CLM2000": "CLIM2000",
    "DOE21E-CIEMAT": "DOE-2.1E CIEMAT", "DOE21E-NREL": "DOE-2.1E NREL", "E+V1": "EnergyPlus",
    "TRN-id": "TRNSYS-ideal", "TRN-re": "TRNSYS-real",
}
REFERENCE_PROGRAMS = list(REFERENCE_SHEETS.values())
CONDENSER_FAN_POWER = 108  # W

CASES = ["CE100", "CE110", "CE120", "CE130", "CE140", "CE150", "CE160", "CE165",
         "CE170", "CE180", "CE185", "CE190", "CE195", "CE200"]

# Output name -> (title, unit, column of the output form, hourly quantity, statistic)
_FORM = [
    ("total_consumption", "Total cooling energy consumption", "kWh", 2, "power", "sum"),
    ("compressor", "Compressor consumption", "kWh", 3, "compressor", "sum"),
    ("indoor_fan", "Supply fan consumption", "kWh", 4, "indoor_fan", "sum"),
    ("condenser_fan", "Condenser fan consumption", "kWh", 5, "condenser_fan", "sum"),
    ("coil_total", "Evaporator coil total load", "kWh", 6, "coil_total", "sum"),
    ("coil_sensible", "Evaporator coil sensible load", "kWh", 7, "coil_sensible", "sum"),
    ("coil_latent", "Evaporator coil latent load", "kWh", 8, "coil_latent", "sum"),
    ("zone_total", "Zone total load", "kWh", 9, "zone_total", "sum"),
    ("zone_sensible", "Zone sensible load", "kWh", 10, "zone_sensible", "sum"),
    ("zone_latent", "Zone latent load", "kWh", 11, "zone_latent", "sum"),
    ("cop_mean", "Mean COP", "-", 12, "cop", "mean"),
    ("idb_mean", "Mean indoor dry bulb", "°C", 13, "zone_temperature", "mean"),
    ("w_mean", "Mean zone humidity ratio", "kg/kg", 14, "zone_humidity_ratio", "mean"),
    ("cop_max", "Maximum COP", "-", 15, "cop", "max"),
    ("idb_max", "Maximum indoor dry bulb", "°C", 16, "zone_temperature", "max"),
    ("w_max", "Maximum zone humidity ratio", "kg/kg", 17, "zone_humidity_ratio", "max"),
    ("cop_min", "Minimum COP", "-", 18, "cop", "min"),
    ("idb_min", "Minimum indoor dry bulb", "°C", 19, "zone_temperature", "min"),
    ("w_min", "Minimum zone humidity ratio", "kg/kg", 20, "zone_humidity_ratio", "min"),
]
OUTPUTS = {name: (f"{title}, February", unit) for name, title, unit, *_ in _FORM}

CELLS = {(case, name, ""): (25 + i, col) for name, _, _, col, _, _ in _FORM for i, case in enumerate(CASES)}

GROUPS = [
    {"title": "Energy consumption (February)", "kind": "annual",
     "outputs": ["total_consumption", "compressor", "indoor_fan", "condenser_fan"]},
    {"title": "Evaporator coil and zone loads (February)", "kind": "annual",
     "outputs": ["coil_total", "coil_sensible", "coil_latent", "zone_total", "zone_sensible", "zone_latent"]},
    {"title": "COP, indoor dry bulb and humidity ratio (February)", "kind": "annual",
     "outputs": ["cop_mean", "cop_max", "cop_min", "idb_mean", "idb_max", "idb_min",
                 "w_mean", "w_max", "w_min"]},
]


def cases():
    return list(CASES)


def case_file(case):
    return FOLDER / "cases" / f"case_{case}.json"


def extract(case, project):
    """Standard's outputs of a simulated case: list of (output, point, value, date) tuples."""
    hours = project.parameter("time_step").value / 3600
    h = dx_hourly(project, CONDENSER_FAN_POWER)
    february = np.array([d.month == 2 for d in project.dates()])
    rows = []
    for name, _, _, _, quantity, statistic in _FORM:
        x = h[quantity][february]
        if statistic == "sum":
            value = x.sum() * hours / 1000  # Wh -> kWh
        else:
            value = {"mean": np.nanmean, "max": np.nanmax, "min": np.nanmin}[statistic](x)
        rows.append((name, "", float(value), ""))
    return rows


def read_reference():
    """Results of the reference programs: {program: {(case, output, point): value}}."""
    wb = openpyxl.load_workbook(RESULTS_FILE, data_only=True)
    reference = {}
    for sheet, program in REFERENCE_SHEETS.items():
        ws = wb[sheet]
        for i, case in enumerate(CASES):  # labelled "CE100" or 100
            label = str(ws.cell(25 + i, 1).value).strip()
            if label not in (case, case[2:]):
                raise ValueError(f"{RESULTS_FILE.name}[{sheet}] row {25 + i}: {label}, expected {case}")
        reference[program] = {
            key: float(value)
            for key, (row, col) in CELLS.items()
            if isinstance(value := ws.cell(row, col).value, (int, float))
        }
    return reference
