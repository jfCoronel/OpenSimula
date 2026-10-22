"""ASHRAE 140-2023 HVAC equipment performance tests for fuel-fired furnaces (HEATING_TEST folder).

Defines the cases to simulate, how the Standard's outputs are extracted from an OpenSimula
project and where each output is in the Standard's output form (sheet YourData of
Std140_HE_Results.xlsx). The reference programs' sheets have the same layout.

Each result is identified by (case, output, point); all outputs here are single values
(point "").
"""

from pathlib import Path

import numpy as np
import openpyxl

NAME = "HEATING_TEST"
TITLE = "Fuel-fired furnace tests (HE100–HE230)"
FOLDER = Path(__file__).resolve().parent.parent / "HEATING_TEST"
RESULTS_FILE = FOLDER / "Std140_HE_Results.xlsx"
# Sheet of the results workbook -> program name
REFERENCE_SHEETS = {"ESP-HOT": "ESP-r/HOT3000", "EnergyPlusV10": "EnergyPlus", "DOE21E": "DOE-2.1E"}
REFERENCE_PROGRAMS = list(REFERENCE_SHEETS.values())

CASES = ["HE100", "HE110", "HE120", "HE130", "HE140", "HE150", "HE160", "HE170", "HE210", "HE220", "HE230"]
FAN_CASES = ["HE150", "HE160", "HE170", "HE210", "HE220", "HE230"]
TEMPERATURE_CASES = ["HE210", "HE220", "HE230"]
DRAFT_FAN_CASES = ["HE170", "HE210", "HE220", "HE230"]
DRAFT_FAN_POWER = 50  # W, runs whenever the furnace burner runs

# Output name -> (title, unit)
OUTPUTS = {
    "furnace_load": ("Total furnace load", "GJ"),
    "furnace_input": ("Total furnace input (fuel)", "GJ"),
    "fan_energy": ("Fan energy (circulating and draft fans)", "kWh"),
    "mean_temperature": ("Mean zone temperature", "°C"),
    "maximum_temperature": ("Maximum zone temperature", "°C"),
    "minimum_temperature": ("Minimum zone temperature", "°C"),
}


def _cells():
    """Position (row, column) of every output in the Standard's output form (column B)."""
    cells = {}
    for output, first_row, cases in [
        ("furnace_load", 20, CASES),
        ("furnace_input", 36, CASES),
        ("fan_energy", 68, FAN_CASES),
        ("mean_temperature", 79, TEMPERATURE_CASES),
        ("maximum_temperature", 87, TEMPERATURE_CASES),
        ("minimum_temperature", 95, TEMPERATURE_CASES),
    ]:
        for i, case in enumerate(cases):
            cells[(case, output, "")] = (first_row + i, 2)
    return cells


CELLS = _cells()

GROUPS = [
    {"title": "Furnace load and fuel input", "kind": "annual", "outputs": ["furnace_load", "furnace_input"]},
    {"title": "Fan energy", "kind": "annual", "outputs": ["fan_energy"]},
    {"title": "Zone temperatures", "kind": "annual",
     "outputs": ["minimum_temperature", "mean_temperature", "maximum_temperature"]},
]


def cases():
    return list(CASES)


def case_file(case):
    return FOLDER / "cases" / f"case_{case}.json"


def extract(case, project):
    """Standard's outputs of a simulated case: list of (output, point, value, date) tuples."""
    seconds = project.parameter("time_step").value
    system = project.component("system")

    def total(variable):  # time integral of a power variable, J
        return system.variable(variable).values.sum() * seconds

    fan = total("indoor_fan_power")
    rows = [
        ("furnace_load", "", total("Q_sensible") / 1e9, ""),
        ("furnace_input", "", (total("power") - fan) / 1e9, ""),
    ]
    if case in FAN_CASES:
        if case in DRAFT_FAN_CASES:  # burner run time fraction of each time step (0 with no load)
            f_load = system.variable("F_load").values
            degradation = system.variable("efficiency_degradation").values
            run_fraction = np.divide(f_load, degradation, out=np.zeros_like(f_load), where=degradation > 0)
            fan += DRAFT_FAN_POWER * run_fraction.sum() * seconds
        rows.append(("fan_energy", "", fan / 3.6e6, ""))
    if case in TEMPERATURE_CASES:
        temp = project.component("spaces_1").variable("temperature").values
        rows.append(("mean_temperature", "", temp.mean(), ""))
        rows.append(("maximum_temperature", "", temp.max(), ""))
        rows.append(("minimum_temperature", "", temp.min(), ""))
    return rows


def read_reference():
    """Example results of the reference programs: {program: {(case, output, point): value}}."""
    wb = openpyxl.load_workbook(RESULTS_FILE, data_only=True)
    reference = {}
    for sheet, program in REFERENCE_SHEETS.items():
        ws = wb[sheet]
        for (case, _, _), (row, _) in CELLS.items():  # rows are labelled "CASE HE100"
            label = str(ws.cell(row, 1).value)
            if case not in label:
                raise ValueError(f"{RESULTS_FILE.name}[{sheet}] row {row}: {label}, expected {case}")
        reference[program] = {
            key: float(value)
            for key, (row, col) in CELLS.items()
            if isinstance(value := ws.cell(row, col).value, (int, float))
        }
    return reference
