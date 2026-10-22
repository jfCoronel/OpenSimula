"""ASHRAE 140-2023 ground-coupled slab-on-grade analytical verification tests GC30b–GC80c
(GROUND_TEST folder, results workbook Std140_GC_Results.xlsx).

Steady-state cases (GC30b, GC30c, GC60b, GC65b) report the floor conduction and zone load at
the end of the simulation; harmonic cases (GC40b–GC80c, two-year simulations) report sums,
means and hourly maxima of the final year, and GC40b/GC40c also the hourly floor conduction.
Each result is identified by (case, output, point): point is "" or the hour of the year.
"""

from pathlib import Path

import numpy as np
import openpyxl

NAME = "GROUND_TEST"
TITLE = "Ground-coupled slab-on-grade tests (GC30b–GC80c)"
FOLDER = Path(__file__).resolve().parent.parent / "GROUND_TEST"
RESULTS_FILE = FOLDER / "Std140_GC_Results.xlsx"
REFERENCE_SHEETS = {"Basecalc": "BASECALC", "EnergyPlus": "EnergyPlus", "ESPr": "ESP-r",
                    "FLUENT": "FLUENT", "GHT": "GHT", "MATLAB": "MATLAB", "SUNREL": "SUNREL-GC",
                    "TRNSYS": "TRNSYS", "VA114": "VA114"}
REFERENCE_PROGRAMS = list(REFERENCE_SHEETS.values())

STEADY_CASES = {"GC30b": 60, "GC30c": 61, "GC60b": 62, "GC65b": 63}  # case -> row
HARMONIC_CASES = {"GC40b": 71, "GC45b": 72, "GC50b": 73, "GC55b": 74, "GC70b": 75, "GC80b": 76,
                  "GC40c": 77, "GC45c": 78, "GC55c": 79, "GC80c": 80}
HOURLY_CASES = {"GC40b": 6, "GC40c": 7}  # case -> column of the hourly floor conduction
HOURS = [str(h) for h in range(1, 8761)]

OUTPUTS = {
    "qfloor": ("Steady-state floor conduction", "W"),
    "qzone": ("Steady-state zone load", "W"),
    "Qfloor": ("Annual floor conduction", "kWh"),
    "Qzone": ("Annual zone heating load", "kWh"),
    "tzone_mean": ("Annual mean zone temperature", "°C"),
    "qfloor_max": ("Annual hourly maximum floor conduction", "Wh/h"),
    "qzone_max": ("Annual hourly maximum zone heating load", "Wh/h"),
    "todb_min": ("Annual minimum outdoor dry bulb", "°C"),
    "qfloor_hourly": ("Hourly floor conduction, final year", "Wh/h"),
}


def _cells():
    cells = {}
    for case, row in STEADY_CASES.items():
        cells[(case, "qfloor", "")] = (row, 5)
        cells[(case, "qzone", "")] = (row, 6)
    for case, row in HARMONIC_CASES.items():
        for output, col in [("Qfloor", 5), ("Qzone", 6), ("tzone_mean", 7), ("qfloor_max", 9),
                            ("qzone_max", 12), ("todb_min", 15)]:
            cells[(case, output, "")] = (row, col)
    for case, col in HOURLY_CASES.items():
        for i, hour in enumerate(HOURS):
            cells[(case, "qfloor_hourly", hour)] = (131 + i, col)
    return cells


CELLS = _cells()

GROUPS = [
    {"title": "Steady-state cases", "kind": "annual", "outputs": ["qfloor", "qzone"]},
    {"title": "Harmonic cases: annual sums and mean", "kind": "annual",
     "outputs": ["Qfloor", "Qzone", "tzone_mean"]},
    {"title": "Harmonic cases: annual hourly maxima and minimum outdoor temperature", "kind": "annual",
     "outputs": ["qfloor_max", "qzone_max", "todb_min"]},
    {"title": "Hourly floor conduction, cases GC40b and GC40c", "kind": "profile", "outputs": ["qfloor_hourly"]},
]


def cases():
    return list(STEADY_CASES) + list(HARMONIC_CASES)


def case_file(case):
    return FOLDER / "cases" / f"case_{case}.json"


def _date(d):
    return f"{d.day:02d}-{d:%b} {d.hour + 1}"


def extract(case, project):
    """Standard's outputs of a simulated case: list of (output, point, value, date) tuples."""
    floor = project.component("floor_wall")
    # Floor conduction into the ground (W) and zone heating load (W)
    q_floor = -floor.variable("q_cd1").values * floor.area
    q_zone = project.component("system").variable("Q_sensible").values
    t_zone = project.component("space_1").variable("temperature").values
    if case in STEADY_CASES:
        return [("qfloor", "", float(q_floor[-1]), ""), ("qzone", "", float(q_zone[-1]), "")]
    # Harmonic cases: final year of the multiyear simulation (hourly time steps)
    last = slice(len(q_floor) - 8760, None)
    dates = project.dates()[last]
    q_floor, q_zone, t_zone = q_floor[last], q_zone[last], t_zone[last]
    t_odb = project.parameter("simulation_file_met").component.variable("temperature").values[last]
    i_floor, i_zone, i_odb = int(np.argmax(q_floor)), int(np.argmax(q_zone)), int(np.argmin(t_odb))
    rows = [
        ("Qfloor", "", float(q_floor.sum() / 1000), ""),
        ("Qzone", "", float(q_zone.sum() / 1000), ""),
        ("tzone_mean", "", float(t_zone.mean()), ""),
        ("qfloor_max", "", float(q_floor[i_floor]), _date(dates[i_floor])),
        ("qzone_max", "", float(q_zone[i_zone]), _date(dates[i_zone])),
        ("todb_min", "", float(t_odb[i_odb]), _date(dates[i_odb])),
    ]
    if case in HOURLY_CASES:
        rows.extend(("qfloor_hourly", hour, float(v), "") for hour, v in zip(HOURS, q_floor))
    return rows


def read_reference():
    """Results of the reference programs: {program: {(case, output, point): value}}."""
    wb = openpyxl.load_workbook(RESULTS_FILE, data_only=True)
    reference = {}
    for sheet, program in REFERENCE_SHEETS.items():
        ws = wb[sheet]
        for case, row in {**STEADY_CASES, **HARMONIC_CASES}.items():
            label = str(ws.cell(row, 4).value).strip()
            if label != case:
                raise ValueError(f"{RESULTS_FILE.name}[{sheet}] row {row}: {label}, expected {case}")
        # Hourly columns only where the sheet uses the output form's layout
        standard_hourly = all(str(ws.cell(128, col).value).strip() == case for case, col in HOURLY_CASES.items())
        reference[program] = {
            key: float(value)
            for key, (row, col) in CELLS.items()
            if (key[1] != "qfloor_hourly" or standard_hourly)
            and isinstance(value := ws.cell(row, col).value, (int, float))
        }
    return reference
