"""ASHRAE 140-2023 HVAC equipment performance tests, unitary space cooling, dynamic cases
CE300–CE545 (COOLING_TEST folder, results workbook Std140_CE-b_Results.xlsx).

Each result is identified by (case, output, point): point is "" for annual values, "May-Sep"
for the May 1 – Sep 30 sums of cases CE500 and CE510, the hour for the June 28 hourly output
of case CE300, and "Apr 30" / "Jun 25" for the daily averages of cases CE500 and CE530.
"""

from pathlib import Path

import numpy as np
import openpyxl

from dx_system import dx_hourly

NAME = "COOLING_TEST_B"
TITLE = "Space cooling equipment, dynamic tests (CE300–CE545)"
FOLDER = Path(__file__).resolve().parent.parent / "COOLING_TEST"
RESULTS_FILE = FOLDER / "Std140_CE-b_Results.xlsx"
REFERENCE_SHEETS = {"TRNSYS-TUD": "TRNSYS-TUD", "DOE22": "DOE-2.2", "DOE21E": "DOE-2.1E",
                    "EnergyPlus1.0": "EnergyPlus", "CodyRun": "CODYRUN", "HOT3000": "HOT3000"}
REFERENCE_PROGRAMS = list(REFERENCE_SHEETS.values())
CONDENSER_FAN_POWER = 930  # W

CASES = ["CE300", "CE310", "CE320", "CE330", "CE340", "CE350", "CE360", "CE400", "CE410",
         "CE420", "CE430", "CE440", "CE500", "CE510", "CE520", "CE522", "CE525", "CE530",
         "CE540", "CE545"]
# Rows of the annual table: CE510 has only its May-Sep sums
ANNUAL_ROWS = [(c, "") for c in CASES[:13]] + [("CE500", "May-Sep"), ("CE510", "May-Sep")] + \
              [(c, "") for c in CASES[14:]]
HOURS = [str(h) for h in range(1, 25)]
DAYS = {"Apr 30": (4, 30), "Jun 25": (6, 25)}

# Annual sums and means (columns B-L): name, title, unit, hourly quantity, statistic
_ANNUAL = [
    ("total_consumption", "Total cooling energy consumption", "kWh", "power", "sum"),
    ("compressor", "Compressor consumption", "kWh", "compressor", "sum"),
    ("condenser_fan", "Condenser fan consumption", "kWh", "condenser_fan", "sum"),
    ("indoor_fan", "Indoor fan consumption", "kWh", "indoor_fan", "sum"),
    ("coil_total", "Evaporator coil total load", "kWh", "coil_total", "sum"),
    ("coil_sensible", "Evaporator coil sensible load", "kWh", "coil_sensible", "sum"),
    ("coil_latent", "Evaporator coil latent load", "kWh", "coil_latent", "sum"),
    ("cop2", "COP2", "-", None, "cop2"),
    ("idb_mean", "Mean indoor dry bulb", "°C", "zone_temperature", "mean"),
    ("w_mean", "Mean zone humidity ratio", "kg/kg", "zone_humidity_ratio", "mean"),
    ("rh_mean", "Mean zone relative humidity", "%", "zone_rh", "mean"),
]
# Annual hourly maxima of consumption and loads (column of the value, Q62..AB81)
_MAXIMA = [
    ("max_power", "Maximum hourly consumption (compressor + both fans)", "Wh", "power", 17),
    ("max_coil_sensible", "Maximum hourly coil sensible load", "Wh", "coil_sensible", 20),
    ("max_coil_latent", "Maximum hourly coil latent load", "Wh", "coil_latent", 23),
    ("max_coil_total", "Maximum hourly coil total load", "Wh", "coil_total", 26),
]
# Annual hourly maxima and minima of COP2 and zone conditions (Q89..AN108)
_EXTREMES = [
    ("max_cop2", "Maximum hourly COP2", "-", "cop2", "max", 17),
    ("min_cop2", "Minimum hourly COP2", "-", "cop2", "min", 20),
    ("max_idb", "Maximum indoor dry bulb", "°C", "zone_temperature", "max", 23),
    ("min_idb", "Minimum indoor dry bulb", "°C", "zone_temperature", "min", 26),
    ("max_w", "Maximum zone humidity ratio", "kg/kg", "zone_humidity_ratio", "max", 29),
    ("min_w", "Minimum zone humidity ratio", "kg/kg", "zone_humidity_ratio", "min", 32),
    ("max_rh", "Maximum zone relative humidity", "%", "zone_rh", "max", 35),
    ("min_rh", "Minimum zone relative humidity", "%", "zone_rh", "min", 38),
]
# Case CE300 June 28 hourly output (B89..L112)
_JUNE28 = [
    ("jun28_compressor", "Compressor consumption", "Wh", "compressor"),
    ("jun28_condenser_fan", "Condenser fan consumption", "Wh", "condenser_fan"),
    ("jun28_coil_total", "Coil total load", "Wh", "coil_total"),
    ("jun28_coil_sensible", "Coil sensible load", "Wh", "coil_sensible"),
    ("jun28_coil_latent", "Coil latent load", "Wh", "coil_latent"),
    ("jun28_w", "Zone humidity ratio", "kg/kg", "zone_humidity_ratio"),
    ("jun28_cop2", "COP2", "-", "cop2"),
    ("jun28_odb", "Outdoor dry bulb", "°C", "outdoor_temperature"),
    ("jun28_edb", "Entering dry bulb", "°C", "entering_drybulb"),
    ("jun28_ewb", "Entering wet bulb", "°C", "entering_wetbulb"),
    ("jun28_outdoor_w", "Outdoor humidity ratio", "kg/kg", "outdoor_humidity_ratio"),
]
# Cases CE500 and CE530 daily average outputs (B..L of rows 120-121 and 129-130)
_DAILY = [
    ("daily_total_consumption", "Total consumption", "Wh", "power", "sum"),
    ("daily_compressor", "Compressor consumption", "Wh", "compressor", "sum"),
    ("daily_condenser_fan", "Condenser fan consumption", "Wh", "condenser_fan", "sum"),
    ("daily_indoor_fan", "Indoor fan consumption", "Wh", "indoor_fan", "sum"),
    ("daily_coil_total", "Coil total load", "Wh", "coil_total", "sum"),
    ("daily_coil_sensible", "Coil sensible load", "Wh", "coil_sensible", "sum"),
    ("daily_coil_latent", "Coil latent load", "Wh", "coil_latent", "sum"),
    ("daily_w", "Zone humidity ratio", "kg/kg", "zone_humidity_ratio", "mean"),
    ("daily_cop2", "COP2", "-", None, "cop2"),
    ("daily_odb", "Outdoor dry bulb", "°C", "outdoor_temperature", "mean"),
    ("daily_edb", "Entering dry bulb", "°C", "entering_drybulb", "mean"),
]

OUTPUTS = {
    **{n: (f"Annual {t[0].lower()}{t[1:]}", u) for n, t, u, *_ in _ANNUAL},
    "outdoor_temperature": ("Annual mean outdoor dry bulb", "°C"),
    "outdoor_humidity_ratio": ("Annual mean outdoor humidity ratio", "kg/kg"),
    **{n: (t, u) for n, t, u, *_ in _MAXIMA},
    "max_odb": ("Maximum outdoor dry bulb", "°C"),
    "max_outdoor_w": ("Maximum outdoor humidity ratio", "kg/kg"),
    **{n: (t, u) for n, t, u, *_ in _EXTREMES},
    **{n: (f"{t}, June 28", u) for n, t, u, _ in _JUNE28},
    **{n: (f"Daily average {t[0].lower()}{t[1:]}", u) for n, t, u, *_ in _DAILY},
}


def _cells():
    cells = {}
    for i, (case, point) in enumerate(ANNUAL_ROWS):
        for col, (name, *_) in enumerate(_ANNUAL, 2):
            cells[(case, name, point)] = (62 + i, col)
    cells[("CE300", "outdoor_temperature", "")] = (62, 13)
    cells[("CE300", "outdoor_humidity_ratio", "")] = (62, 14)
    for i, case in enumerate(CASES):
        for name, *_, col in _MAXIMA:
            cells[(case, name, "")] = (62 + i, col)
        for name, *_, col in _EXTREMES:
            cells[(case, name, "")] = (89 + i, col)
    cells[("CE300", "max_odb", "")] = (62, 29)
    cells[("CE300", "max_outdoor_w", "")] = (62, 32)
    for col, (name, *_) in enumerate(_JUNE28, 2):
        for h, hour in enumerate(HOURS):
            cells[("CE300", name, hour)] = (89 + h, col)
    for case, first_row in [("CE500", 120), ("CE530", 129)]:
        for d, day in enumerate(DAYS):
            for col, (name, *_) in enumerate(_DAILY, 2):
                cells[(case, name, day)] = (first_row + d, col)
    return cells


CELLS = _cells()

GROUPS = [
    {"title": "Annual energy consumption", "kind": "annual",
     "outputs": ["total_consumption", "compressor", "condenser_fan", "indoor_fan"]},
    {"title": "Annual evaporator coil loads and COP2", "kind": "annual",
     "outputs": ["coil_total", "coil_sensible", "coil_latent", "cop2"]},
    {"title": "Annual mean zone conditions", "kind": "annual",
     "outputs": ["idb_mean", "w_mean", "rh_mean"]},
    {"title": "Annual hourly maxima of consumption and loads", "kind": "annual",
     "outputs": [n for n, *_ in _MAXIMA]},
    {"title": "Annual hourly maxima and minima of COP2 and zone conditions", "kind": "annual",
     "outputs": [n for n, *_ in _EXTREMES]},
    {"title": "Weather checks, case CE300", "kind": "annual",
     "outputs": ["outdoor_temperature", "outdoor_humidity_ratio", "max_odb", "max_outdoor_w"]},
    {"title": "June 28 hourly output, case CE300", "kind": "profile", "outputs": [n for n, *_ in _JUNE28]},
    {"title": "Daily average outputs, cases CE500 and CE530", "kind": "annual",
     "outputs": [n for n, *_ in _DAILY]},
]


def cases():
    return list(CASES)


def case_file(case):
    return FOLDER / "cases" / f"case_{case}.json"


def _date(d):
    """Day, month and hour (1-24, hour ending) of a time step."""
    return f"{d.day:02d}-{d:%b} {d.hour + 1}"


def _cop2(h, mask):
    work = h["compressor"][mask].sum() + h["condenser_fan"][mask].sum()
    return h["coil_total"][mask].sum() / work if work > 0 else float("nan")


def extract(case, project):
    """Standard's outputs of a simulated case: list of (output, point, value, date) tuples."""
    hours = project.parameter("time_step").value / 3600
    dates = project.dates()
    h = dx_hourly(project, CONDENSER_FAN_POWER)
    wanted = {(o, p) for c, o, p in CELLS if c == case}
    rows = []

    def period(start, end):  # inclusive dates (month, day)
        return np.array([start <= (d.month, d.day) <= end for d in dates])

    for point, mask in [("", np.ones(len(dates), bool)), ("May-Sep", period((5, 1), (9, 30)))]:
        if ("total_consumption", point) not in wanted:
            continue
        for name, _, _, quantity, statistic in _ANNUAL:
            if statistic == "sum":
                value = h[quantity][mask].sum() * hours / 1000
            elif statistic == "mean":
                value = h[quantity][mask].mean()
            else:
                value = _cop2(h, mask)
            rows.append((name, point, float(value), ""))
    if ("outdoor_temperature", "") in wanted:
        rows.append(("outdoor_temperature", "", float(h["outdoor_temperature"].mean()), ""))
        rows.append(("outdoor_humidity_ratio", "", float(h["outdoor_humidity_ratio"].mean()), ""))
        for name, quantity in [("max_odb", "outdoor_temperature"), ("max_outdoor_w", "outdoor_humidity_ratio")]:
            i = int(np.argmax(h[quantity]))
            rows.append((name, "", float(h[quantity][i]), _date(dates[i])))
    for name, _, _, quantity, _ in _MAXIMA:
        i = int(np.argmax(h[quantity]))
        rows.append((name, "", float(h[quantity][i] * hours), _date(dates[i])))
    # For CE500-CE545 the minima of IDB, humidity ratio and RH use April 1 - December 31
    late = period((4, 1), (12, 31)) if case >= "CE500" else np.ones(len(dates), bool)
    for name, _, _, quantity, statistic, _ in _EXTREMES:
        mask = late if statistic == "min" and quantity != "cop2" else np.ones(len(dates), bool)
        x = np.where(mask, h[quantity], np.nan)
        i = int(np.nanargmax(x) if statistic == "max" else np.nanargmin(x))
        rows.append((name, "", float(x[i]), _date(dates[i])))
    if ("jun28_compressor", "1") in wanted:
        day = period((6, 28), (6, 28))
        for name, _, unit, quantity in _JUNE28:
            factor = hours if unit == "Wh" else 1
            rows.extend((name, hour, float(v * factor), "") for hour, v in zip(HOURS, h[quantity][day]))
    for day, (m, d) in DAYS.items():
        if ("daily_total_consumption", day) not in wanted:
            continue
        mask = period((m, d), (m, d))
        for name, _, _, quantity, statistic in _DAILY:
            if statistic == "sum":  # average hourly value of the day: mean W x 1 h = Wh
                value = h[quantity][mask].mean()
            elif statistic == "mean":
                value = h[quantity][mask].mean()
            else:
                value = _cop2(h, mask)
            rows.append((name, day, float(value), ""))
    return rows


def read_reference():
    """Results of the reference programs: {program: {(case, output, point): value}}."""
    wb = openpyxl.load_workbook(RESULTS_FILE, data_only=True)
    reference = {}
    for sheet, program in REFERENCE_SHEETS.items():
        ws = wb[sheet]
        checks = [(62 + i, 1, case) for i, (case, _) in enumerate(ANNUAL_ROWS)] + \
                 [(62 + i, 16, case) for i, case in enumerate(CASES)] + \
                 [(89 + i, 16, case) for i, case in enumerate(CASES)]
        for row, col, case in checks:  # labelled "CE300" or "E300"
            label = str(ws.cell(row, col).value).strip()
            if not label.startswith((case, case[1:])):
                raise ValueError(f"{RESULTS_FILE.name}[{sheet}] row {row}: {label}, expected {case}")
        reference[program] = {
            key: float(value)
            for key, (row, col) in CELLS.items()
            if isinstance(value := ws.cell(row, col).value, (int, float))
        }
    return reference
